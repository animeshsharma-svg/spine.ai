"""Style an exported bug workbook like the hand-made original (fonts, header
colours, borders, severity colours, frozen panes, Excel tables).

    python3 scripts/style-bug-sheet.py <file.xlsx>

Called by scripts/export-bug-sheet.ts. Content is never changed here — only
how it looks. Needs openpyxl; the exporter carries on unstyled without it.
"""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

path = sys.argv[1]
wb = load_workbook(path)
F = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
body_font = Font(name=F, size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
sev_fill = {"Critical": "C00000", "High": "F4B084", "Medium": "FFE699", "Low": "E2EFDA"}

def table(ws, name):
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, Alignment(wrap_text=True, vertical="center"), border
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = body_font, wrap, border
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 30
    if ws.max_row >= 2:
        t = Table(displayName=name, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)

for sheet, name in (("Bugs", "Bugs"), ("Needs confirmation", "Confirm"), ("Data problems", "DataProblems")):
    if sheet in wb.sheetnames:
        table(wb[sheet], name)
if "Bugs" in wb.sheetnames:
    ws = wb["Bugs"]
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=9)
        if c.value in sev_fill:
            c.fill = PatternFill("solid", fgColor=sev_fill[c.value])
            if c.value == "Critical":
                c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
if "Summary" in wb.sheetnames:
    s = wb["Summary"]
    s["A1"].font = Font(name=F, bold=True, size=14)
    s["A2"].font = Font(name=F, size=10, italic=True)
    s["A2"].alignment = Alignment(wrap_text=True)
    s.row_dimensions[2].height = 45
    for row in s.iter_rows(min_row=4):
        for c in row:
            if c.value is None:
                continue
            heading = isinstance(c.value, str) and c.column == 1 and s.cell(row=c.row, column=2).value is None
            c.font = Font(name=F, bold=True) if heading else Font(name=F, size=10)
    s["A4"].font = Font(name=F, bold=True, size=11)
    s["B4"].font = Font(name=F, bold=True, size=11)
wb.save(path)

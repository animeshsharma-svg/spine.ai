import json, re
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

C = json.load(open("cases.json"))
COLS = ["Test Case ID", "Module", "Scenario", "Test Case Title", "User Role", "Portal", "Pre-conditions", "Test Data", "Test Steps", "Expected Result", "Priority", "Test Type", "Status"]
WID = [14, 30, 28, 46, 20, 18, 40, 30, 80, 70, 10, 11, 10]
F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1F3A5F"); thin = Side(style="thin", color="D0D5DD")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
PRI = {"Critical": "FDE2E1", "High": "FFF1D6", "Medium": "E8F4EA", "Low": "EEF0F3"}

wb = Workbook(); ws = wb.active; ws.title = "Sourcing Test Cases"
ws.append(COLS)
for i, w in enumerate(WID, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(row=1, column=i); c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BRD
for r, case in enumerate(C, 2):
    for i, k in enumerate(COLS, 1):
        cell = ws.cell(row=r, column=i, value=case[k]); cell.font = F; cell.alignment = WRAP; cell.border = BRD
    ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor=PRI.get(case["Priority"], "FFFFFF"))
ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:M{len(C) + 1}"; ws.row_dimensions[1].height = 30

# summary
s2 = wb.create_sheet("Summary")
s2.append(["Module", "Cases", "Positive", "Negative", "Edge", "Critical", "High"])
for i, w in enumerate([70, 8, 9, 9, 7, 8, 7], 1): s2.column_dimensions[get_column_letter(i)].width = w
mods = []
for c in C:
    if c["Module"] not in mods: mods.append(c["Module"])
for m in mods:
    cs = [c for c in C if c["Module"] == m]
    s2.append([m, len(cs)] + [sum(c["Test Type"] == t for c in cs) for t in ("Positive", "Negative", "Edge")] + [sum(c["Priority"] == p for c in cs) for p in ("Critical", "High")])
s2.append(["Total", len(C)] + [sum(c["Test Type"] == t for c in C) for t in ("Positive", "Negative", "Edge")] + [sum(c["Priority"] == p for c in C) for p in ("Critical", "High")])
for row in s2.iter_rows():
    for cell in row: cell.font = F; cell.border = BRD; cell.alignment = WRAP
for cell in s2[1]: cell.font = FB; cell.fill = HEAD
for cell in s2[s2.max_row]: cell.font = Font(name="Arial", size=10, bold=True)

# product rules quoted in the cases → where they come from
app = json.load(open("../../../procol-ui-tests/src/oracle/appRules.json"))
amap = {}
for r in app["rules"]:
    m = (r.get("message") or "").strip()
    if m: amap.setdefault(m.lower(), r.get("file") or r.get("source") or "")
codes = json.load(open("../../../prototype/data/app-error-catalog.json"))["codes"]
cmap = {}
for code, it in (codes.items() if isinstance(codes, dict) else [(x.get("code"), x) for x in codes]):
    msg = (it.get("message") if isinstance(it, dict) else str(it)) or ""
    if msg: cmap.setdefault(msg.strip().lower(), f"backend error {code}")
EXTRA = {  # read from the dashboard code / translations while writing the cases
    "disabled during bidding time.": "dashboard en.json common.rejectBidOptionDisabled (layerDetails.jsx)",
    "auction cannot be withdrawn when the event is live": "dashboard eventHeader.jsx (isWithdrawDisabledDuringSubmissionTime)",
    "order withdrawn": "dashboard en.json successMessages.orderWithDrawn",
    "are you sure you want to withdraw this event?": "dashboard en.json orderDetails.WithDrawReasonModal",
    "please enter reason for withdrawing this event": "dashboard en.json orderDetails.WithDrawReasonModal.label",
    "supplier would not be allowed to bid again. this action cannot be undone.": "dashboard en.json orderDetails.rejectBidModal",
    "you have blocked the supplier from bidding again.": "dashboard en.json orderDetails.rejectBidModal",
    "no bids rejected for this event": "dashboard en.json eventDetails.rejectBidModal.empty",
    "vendors who have placed the bid, have accepted the terms and conditions.": "dashboard en.json eventDetails.isTncMandatory",
    "terms and conditions are not added": "dashboard en.json eventDetails.tncNotAddedWarning",
    "overridden by mark all mandatory": "dashboard en.json eventCreation.tncModal",
    "bulk counter offer sent successfully!": "dashboard en.json successMessages.bulkCounterOffer",
    "sealed bid time extended!": "dashboard en.json orderDetails.increaseAuctionModal",
    "dormant locked pr cannot be purchased": "dashboard purchaseRequisitions.jsx",
    "rejected or send back pending pr cannot be purchased": "dashboard purchaseRequisitions.jsx",
    "multiple boq-based pr cannot be purchased together": "dashboard purchaseRequisitions.jsx",
    "services and materials cannot be purchased together": "dashboard purchaseRequisitions.jsx",
    "unlock request sent successfully": "dashboard purchaseRequisitions.jsx",
    "line items can only be reopened if at least one line item has expired": "dashboard eventHeader.jsx",
    "you cannot edit the config in traffic event": "dashboard en.json eventDetails.toolTipMessages.trafficLight",
    "no change in audience": "dashboard en.json errorMessages.updateAudience",
}
uses = defaultdict(list)
for c in C:
    for q in re.findall(r"“([^”]{6,200})”|\"([^\"]{6,200})\"", c["Expected Result"]):
        msg = (q[0] or q[1]).strip()
        uses[msg].append(c["Test Case ID"])
s3 = wb.create_sheet("Product rules")
s3.append(["App message quoted in the expected results", "Where it comes from", "Used in test cases"])
for i, w in enumerate([70, 55, 60], 1): s3.column_dimensions[get_column_letter(i)].width = w
found = 0
for msg, ids in sorted(uses.items(), key=lambda kv: -len(kv[1])):
    k = msg.lower().rstrip(".")
    src = amap.get(msg.lower()) or amap.get(k) or cmap.get(msg.lower()) or cmap.get(k) or cmap.get(k + ".") or amap.get(k + ".") or EXTRA.get(msg.lower()) or EXTRA.get(k) or EXTRA.get(k + ".")
    if not src: continue
    found += 1
    s3.append([msg, src if "backend" in src or "dashboard" in src else f"dashboard {src}", ", ".join(ids[:25]) + (f" … (+{len(ids) - 25})" if len(ids) > 25 else "")])
for row in s3.iter_rows():
    for cell in row: cell.font = F; cell.border = BRD; cell.alignment = WRAP
for cell in s3[1]: cell.font = FB; cell.fill = HEAD
s3.freeze_panes = "A2"

out = "../Procol-Sourcing-Test-Cases-PR-to-PO.xlsx"
wb.save(out)
print("saved", out, "| cases", len(C), "| rules traced", found, "of", len(uses), "quoted strings")
